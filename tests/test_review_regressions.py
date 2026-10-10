"""Offline regressions for scan completeness, API shapes, and report boundaries."""
import csv
import importlib.util
import json
import logging
import subprocess
import sys
import tempfile
import unittest
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import Mock, patch

from common import session  # inject corporate CA trust before importing boto3
import boto3
from botocore.exceptions import ClientError
from botocore.stub import Stubber
from openpyxl import Workbook, load_workbook

from common import cli, findings_io, output
import build_dataset
import generate_dashboard
import generate_report

ROOT = Path(__file__).resolve().parents[1]
ACCOUNT = '123456789012'


def load(relative):
    name = Path(relative).stem
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


roles = load('wave4_security/37_detect_iam_risky_roles.py')
iops = load('wave5_performance/40_detect_idle_provisioned_iops.py')
nat = load('wave1_cost/06_detect_nat_gateway_hotpaths.py')
runbook = load('wave3_reliability/28_dr_runbook_tracker.py')
lam = load('wave2_lifecycle/15_detect_lambda_deprecated_runtimes.py')
athena = load('wave1_cost/10_detect_athena_result_buckets.py')
quotas = load('continuous/49_detect_custom_quotas.py')
tagging = load('continuous/43_detect_tagging_compliance.py')
inventory = load('continuous/54_inventory_refresh.py')
compare = load('wave4_security/compare_sg_findings.py')


def client(service):
    return boto3.client(service, region_name='us-east-1',
                        aws_access_key_id='testing', aws_secret_access_key='testing')


def raw_row(**overrides):
    return dict(account_id=ACCOUNT, region='us-east-1', resource_id='resource-demo',
                resource_type='Security Service', finding='guardduty-not-enabled',
                detail='Offline fixture', estimated_monthly_savings_usd=0, owner='ops',
                **overrides)


class ScanPublicationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.name = '30_detect_open_security_groups'
        self.folder = Path(self.temp.name) / 'wave4_security'
        self.folder.mkdir()
        self.old = self.folder / f'{self.name}_20250101_120000.csv'
        self.old.write_text(f'account_id,finding\n{ACCOUNT},open-to-internet\n')

    def test_failure_does_not_retire_existing_findings(self):
        for workers in (1, 2):
            with self.subTest(workers=workers):
                with self.assertRaises(session.IncompleteScanError):
                    with output.FindingsWriter(self.temp.name, self.name) as writer:
                        session.run_parallel([(None, ACCOUNT, 'us-east-1', writer)],
                                             lambda *args: 1 / 0, max_workers=workers)
                self.assertEqual(findings_io.discover_latest_csvs(self.temp.name), [self.old])
        partials = list(self.folder.glob('*.incomplete.csv'))
        self.assertEqual(len(partials), 2)
        manifest = json.loads(partials[0].with_suffix('.json').read_text())
        self.assertEqual(manifest['status'], 'incomplete')
        self.assertFalse(manifest['tasks'][0]['complete'])

    def test_successful_empty_scan_retires_old_findings(self):
        with output.FindingsWriter(self.temp.name, self.name) as writer:
            outcomes = session.run_parallel([(None, ACCOUNT)], lambda *args: None)
            writer.record_outcomes(outcomes)
        latest = findings_io.discover_latest_csvs(self.temp.name)
        self.assertEqual(latest, [writer._path])
        with writer._path.open() as fh:
            self.assertEqual(list(csv.DictReader(fh)), [])
        self.assertTrue(json.loads(writer._path.with_suffix('.json').read_text())['tasks'][0]['complete'])

    def test_sample_and_narrower_scans_stay_out_of_all_consumers(self):
        for kwargs in ({'sample': True}, {'scope': {'accounts': ['234567890123']}}):
            with output.FindingsWriter(self.temp.name, self.name, **kwargs):
                pass
        self.assertEqual(generate_report.discover_csvs(self.temp.name), [self.old])
        self.assertEqual(inventory._collect_csvs(Path(self.temp.name), None), [self.old])
        self.assertEqual(compare._find_latest_our_csv(Path(self.temp.name)), self.old)

    def test_writer_reserves_distinct_paths_within_one_second(self):
        with output.FindingsWriter(self.temp.name, self.name) as a:
            with output.FindingsWriter(self.temp.name, self.name) as b:
                self.assertNotEqual(a._path, b._path)
        self.assertTrue(a._path.exists())
        self.assertTrue(b._path.exists())

    def test_readers_reject_corrupt_manifest(self):
        self.old.with_suffix('.json').write_text('{bad json')
        self.assertEqual(findings_io.discover_latest_csvs(self.temp.name), [])

    def test_skipped_access_and_logged_errors_fail_tasks(self):
        def skip(*args):
            session.is_skippable(ClientError({'Error': {'Code': 'AccessDenied'}}, 'ListResources'))
        def log_only(*args):
            logging.getLogger('detector').error('Enumeration failed')
        for fn in (skip, log_only):
            with self.subTest(fn=fn.__name__), self.assertRaises(session.IncompleteScanError):
                session.run_parallel([(None, ACCOUNT)], fn)

    def test_all_tasks_run_and_thread_results_are_isolated(self):
        seen = []
        def scan(_, account):
            seen.append(account)
            if account == 'bad':
                session.mark_incomplete('missing metrics')
        with self.assertRaises(session.IncompleteScanError) as caught:
            session.run_parallel([(None, 'bad'), (None, 'good')], scan, 2)
        outcomes = {o.target: o.complete for o in caught.exception.outcomes}
        self.assertEqual(outcomes, {'bad': False, 'good': True})
        self.assertCountEqual(seen, ['bad', 'good'])

    def test_missing_credentials_and_region_discovery_do_not_default_clean(self):
        with patch.object(session, 'BritiveKeychain') as kc:
            kc.return_value.account_keys = {}
            with self.assertRaises(session.IncompleteScanError):
                with session.AccountIterator([ACCOUNT]):
                    pass
        sess = Mock()
        sess.client.return_value.describe_regions.side_effect = RuntimeError('No access')
        with self.assertRaises(session.IncompleteScanError):
            session._get_enabled_regions(sess)
    def test_cli_marks_explicit_region_selection_as_sample(self):
        args = cli.make_parser('test').parse_args([
            '--accounts', ACCOUNT, '--output-dir', self.temp.name,
            '--regions', 'us-east-1', '--workers', '1'])
        with patch.object(session, 'AccountIterator') as iterator:
            iterator.return_value.__enter__.return_value.iter_regions.return_value = [(None, ACCOUNT, 'us-east-1')]
            cli.run(args, self.name, lambda *args: None)
        self.assertEqual(findings_io.discover_latest_csvs(self.temp.name), [self.old])
        samples = list(self.folder.glob('*.sample.csv'))
        self.assertEqual(len(samples), 1)
        manifest = json.loads(samples[0].with_suffix('.json').read_text())
        self.assertEqual(manifest['scope']['accounts'], [ACCOUNT])
        self.assertEqual(manifest['scope']['regions'], ['us-east-1'])

    def test_failed_run_does_not_prune_successful_history(self):
        with self.assertRaises(ValueError):
            with output.FindingsWriter(self.temp.name, self.name, keep_runs=1):
                raise ValueError('failed sweep')
        self.assertTrue(self.old.exists())



class DetectorTests(unittest.TestCase):
    def test_iam_fetches_get_role_for_active_role(self):
        iam = client('iam')
        now = datetime.now(timezone.utc)
        listed = dict(Path='/', RoleName='active-role', RoleId='AROATESTROLE12345',
                      Arn=f'arn:aws:iam::{ACCOUNT}:role/active-role', CreateDate=now-timedelta(days=400))
        with Stubber(iam) as stub:
            stub.add_response('list_roles', {'Roles': [listed], 'IsTruncated': False}, {})
            stub.add_response('get_role', {'Role': {**listed,
                'RoleLastUsed': {'LastUsedDate': now-timedelta(days=1), 'Region': 'us-east-1'}}},
                {'RoleName': 'active-role'})
            sess = Mock()
            sess.client.return_value = iam
            writer = Mock()
            with patch.object(roles, '_check_attached_policies'), patch.object(roles, '_check_inline_policies'):
                roles.scan(sess, ACCOUNT, writer)
            writer.write.assert_not_called()
            stub.assert_no_pending_responses()

    def test_old_role_without_usage_is_not_called_never_used(self):
        now = datetime.now(timezone.utc)
        role = dict(RoleName='old-role', Arn=f'arn:aws:iam::{ACCOUNT}:role/old-role',
                    CreateDate=now-timedelta(days=400))
        iam = Mock()
        iam.get_paginator.return_value.paginate.return_value = [{'Roles': [role]}]
        iam.get_role.return_value = {'Role': role}
        sess = Mock()
        sess.client.return_value = iam
        writer = Mock()
        with patch.object(roles, '_check_attached_policies'), patch.object(roles, '_check_inline_policies'):
            roles.scan(sess, ACCOUNT, writer)
        self.assertEqual(writer.write.call_args.kwargs['finding'], 'role-no-recorded-use')

    def test_missing_or_partial_iops_is_unknown_but_measured_zero_is_valid(self):
        for values, status, expected in (([], 'Complete', None), ([0], 'PartialData', None),
                                         ([0], 'Complete', 0.0)):
            cw = Mock()
            cw.get_paginator.return_value.paginate.return_value = [{'MetricDataResults': [
                {'Id': f'v0m{i}', 'Values': values, 'StatusCode': status} for i in range(2)]}]
            with self.subTest(values=values, status=status):
                self.assertEqual(iops._batch_iops(cw, ['vol-test']), {'vol-test': expected})
        cw.get_paginator.return_value.paginate.side_effect = ClientError(
            {'Error': {'Code': 'AccessDenied'}}, 'GetMetricData')
        self.assertEqual(iops._batch_iops(cw, ['vol-test']), {'vol-test': None})

    def test_nat_missing_data_is_not_an_idle_gateway(self):
        cw = Mock()
        cw.get_paginator.return_value.paginate.return_value = [{'MetricDataResults': [
            {'Id': 'q0', 'Values': [], 'StatusCode': 'Complete'}]}]
        self.assertEqual(nat._batch_bytes(cw, ['nat-test']), {})
        cw.get_paginator.return_value.paginate.return_value[0]['MetricDataResults'][0]['Values'] = [0]
        self.assertEqual(nat._batch_bytes(cw, ['nat-test']), {'nat-test': 0.0})

    def test_runbook_valid_dates_and_empty_snapshot(self):
        with tempfile.TemporaryDirectory() as d:
            registry = Path(d) / 'registry.csv'
            registry.write_text('workload_name,last_test_date\nold,2020-01-01\n')
            self.assertEqual(len(runbook.check_registry(registry, 180)), 1)
            registry.write_text(f'workload_name,last_test_date\nrecent,{date.today().isoformat()}\n')
            self.assertEqual(runbook.check_registry(registry, 180), [])
            path = runbook.write_findings([], Path(d)/'findings')
            self.assertEqual(findings_io.discover_latest_csvs(Path(d)/'findings'), [path])

    def test_lambda_deprecated_and_upcoming_runtimes(self):
        for runtime in ('ruby3.2', 'dotnet6', 'provided.al2', 'python3.8'):
            self.assertEqual(lam._classify(runtime, date(2026, 10, 6))[1], 'deprecated-runtime')
        self.assertEqual(lam._classify('python3.10', date(2026, 10, 6))[1], 'runtime-eol-soon')
        self.assertIsNone(lam._classify('python3.12', date(2026, 10, 6)))

    def test_lifecycle_filters_cannot_prove_bucket_wide_coverage(self):
        for filter_data in ({'Prefix': 'one/'}, {'Tag': {'Key': 'app', 'Value': 'one'}}, {}):
            s3 = Mock()
            s3.get_bucket_lifecycle_configuration.return_value = {'Rules': [{
                'Status': 'Enabled', 'Filter': filter_data, 'Expiration': {'Days': 7},
                'AbortIncompleteMultipartUpload': {'DaysAfterInitiation': 7}}]}
            expected = (bool(filter_data), bool(filter_data))
            self.assertEqual(athena._get_lifecycle_gaps(s3, 'bucket'), expected)

    def test_lambda_custom_quota_uses_requested_region(self):
        sess = Mock()
        sess.client.return_value.get_paginator.return_value.paginate.return_value = []
        quotas._check_lambda_vpc_eni_count(sess, 'eu-west-1')
        sess.client.assert_called_once_with('lambda', region_name='eu-west-1')

    def test_athena_lifecycle_access_failure_does_not_publish_complete_scan(self):
        sess = Mock()
        s3 = sess.client.return_value
        s3.list_buckets.return_value = {'Buckets': [{'Name': 'aws-athena-query-results-test'}]}
        s3.get_bucket_lifecycle_configuration.side_effect = ClientError(
            {'Error': {'Code': 'AccessDenied'}}, 'GetBucketLifecycleConfiguration')
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaises(session.IncompleteScanError):
                with output.FindingsWriter(folder, athena.SCRIPT_NAME) as writer:
                    session.run_parallel([(sess, ACCOUNT, writer)], athena.scan)
            self.assertEqual(findings_io.discover_latest_csvs(folder), [])
            self.assertEqual(len(list(Path(folder).rglob('*.incomplete.csv'))), 1)

    def test_tagging_merges_untagged_resources_and_derives_type(self):
        tg = Mock()
        arn = f'arn:aws:ec2:us-east-1:{ACCOUNT}:instance/i-test'
        tg.get_paginator.return_value.paginate.return_value = [
            {'ResourceTagMappingList': [{'ResourceARN': arn, 'Tags': []}]}]
        ex = Mock()
        ex.get_default_view.return_value = {'ViewArn': 'view'}
        ex.get_view.return_value = {'View': {'IncludedProperties': [{'Name': 'tags'}]}}
        ex.search.side_effect = [
            {'Count': {'Complete': True}, 'Resources': [{'Arn': arn}], 'NextToken': 'page2'},
            {'Count': {'Complete': True}, 'Resources': [{'Arn': arn.replace('i-test','i-new')}]}]
        resources = list(tagging._iter_resources(tg, ex, ['ec2:instance'], 'us-east-1'))
        self.assertEqual(len(resources), 2)
        self.assertEqual(tagging._resource_type_from_arn(arn), 'ec2:instance')
        self.assertEqual(ex.search.call_args.kwargs['NextToken'], 'page2')
        self.assertIn('tag:none', ex.search.call_args.kwargs['QueryString'])

    def test_resource_explorer_requests_match_sdk_models(self):
        tg, ex = client('resourcegroupstaggingapi'), client('resource-explorer-2')
        arn = f'arn:aws:ec2:us-east-1:{ACCOUNT}:instance/i-new'
        view = f'arn:aws:resource-explorer-2:us-east-1:{ACCOUNT}:view/a/1234'
        query = 'tag:none -tag.key:aws* resourcetype.supports:tags resourcetype:ec2:instance region:us-east-1'
        with Stubber(tg) as tags_stub, Stubber(ex) as explorer_stub:
            explorer_stub.add_response('get_default_view', {'ViewArn': view}, {})
            explorer_stub.add_response('get_view', {'View': {
                'ViewArn': view, 'IncludedProperties': [{'Name': 'tags'}],
                'Filters': {'FilterString': ''}}}, {'ViewArn': view})
            tags_stub.add_response('get_resources', {'ResourceTagMappingList': []},
                                   {'ResourceTypeFilters': ['ec2:instance']})
            explorer_stub.add_response('search', {'Resources': [{'Arn': arn}],
                'Count': {'Complete': True, 'TotalResources': 1}, 'ViewArn': view},
                {'QueryString': query, 'ViewArn': view, 'MaxResults': 100})
            resources = list(tagging._iter_resources(tg, ex, ['ec2:instance'], 'us-east-1'))
            self.assertEqual(resources, [{'ResourceARN': arn, 'Tags': []}])
            tags_stub.assert_no_pending_responses()
            explorer_stub.assert_no_pending_responses()

    def test_tagging_rejects_filtered_view_and_truncated_search(self):
        tg = Mock()
        tg.get_paginator.return_value.paginate.return_value = []
        ex = Mock()
        ex.get_default_view.return_value = {'ViewArn': 'view'}
        for view, response in (({'IncludedProperties': []}, {}),
                               ({'IncludedProperties': [{'Name':'tags'}], 'Filters': {'FilterString':'service:ec2'}}, {}),
                               ({'IncludedProperties': [{'Name':'tags'}]}, {'Count': {'Complete': False}})):
            ex.get_view.return_value = {'View': view}
            ex.search.return_value = response
            with self.assertRaises(session.IncompleteScanError):
                list(tagging._iter_resources(tg, ex, ['ec2:instance'], 'us-east-1'))


class ReportingTests(unittest.TestCase):
    def test_severity_matches_between_dataset_and_excel(self):
        cases = [('wave4_security', 'guardduty-not-enabled', 'High', generate_report.rag_security),
                 ('wave4_security', 'securityhub-not-enabled', 'High', generate_report.rag_security),
                 ('wave4_security', 'config-recorder-stopped', 'High', generate_report.rag_security),
                 ('wave4_security', 'console-no-mfa', 'High', generate_report.rag_security),
                 ('wave3_reliability', 'dr-quota-lower-than-primary', 'High', generate_report.rag_reliability),
                 ('wave3_reliability', 'untested-runbook', 'Medium', generate_report.rag_reliability),
                 ('wave4_security', 'icmp-echo-open-to-internet', 'Low', generate_report.rag_security)]
        fills = {'High': generate_report._RED_FILL, 'Medium': generate_report._AMBER_FILL,
                 'Low': generate_report._GREEN_FILL}
        for wave, finding, expected, rag in cases:
            with self.subTest(finding=finding):
                self.assertEqual(findings_io.severity_of(wave, finding, 0), expected)
                self.assertEqual(rag(finding, 0)[0], fills[expected])

    def test_excel_owner_is_literal_on_all_data_sheets(self):
        row = build_dataset.normalize_row(raw_row(), '36_detect_security_services_gaps', '2026-10-06')
        row.update(owner='=1+1', estimated_monthly_savings_usd=100,
                   estimated_annual_savings_usd=1200, resource_type='=1+2', detail='=HYPERLINK("https://example.invalid", "text")')
        wb = Workbook()
        generate_report.write_all_findings_sheet(wb, [row])
        generate_report.write_wave_sheet(wb, [row], 'wave4_security', 'Security', generate_report.rag_security)
        generate_report.write_top_savings_sheet(wb, [row])
        generate_report.write_summary_sheet(wb, [row])
        generate_report.write_by_account_sheet(wb, [row])
        with tempfile.TemporaryDirectory() as d:
            path = Path(d)/'report.xlsx'
            wb.save(path)
            reloaded = load_workbook(path)
            for ws in reloaded:
                for cells in ws.iter_rows():
                    for cell in cells:
                        if isinstance(cell.value, str) and cell.value.startswith('='):
                            self.assertEqual(cell.data_type, 's')

    def test_dashboard_escapes_static_and_dynamic_cells(self):
        row = build_dataset.normalize_row(raw_row(), '36_detect_security_services_gaps', '2026-10-06')
        payload = '<img src=x onerror=alert(1)>'
        row.update(resource_type=payload, owner=payload)
        html = generate_dashboard.build_html([row], [row], '')
        self.assertNotIn('<td>' + payload + '</td>', html)
        self.assertIn('&lt;img src=x onerror=alert(1)&gt;', html)
        self.assertIn('esc(r.owner || "-")', html)
        self.assertNotIn(payload, generate_dashboard._svg_bar([{'name': payload, 'value': 1}], 'name','value','chart'))

    def test_dataset_keeps_observation_time(self):
        row = raw_row()
        row['scan_time'] = '2026-09-01T12:00:00+00:00'
        normalized = build_dataset.normalize_row(row, '36_detect_security_services_gaps', '2026-10-06')
        self.assertEqual(normalized['scan_time'], row['scan_time'])

    def test_offline_reporting_cli_pipeline(self):
        with tempfile.TemporaryDirectory() as d:
            work = Path(d)
            findings = work/'findings'
            with output.FindingsWriter(str(findings), '36_detect_security_services_gaps') as writer:
                writer.write(**raw_row())
            commands = [
                ['build_dataset.py','--findings-dir',str(findings),'--dataset-file',str(work/'dataset.csv')],
                ['continuous/54_inventory_refresh.py','--findings-dir',str(findings),'--output-dir',str(findings)],
                ['generate_report.py','--findings-dir',str(findings),'--output-dir',str(work/'reports')],
                ['generate_report.py','--from-inventory','--findings-dir',str(findings),'--output-dir',str(work/'inventory_reports')],
                ['generate_dashboard.py','--dataset-file',str(work/'dataset.csv'),'--output-dir',str(work/'reports')],
            ]
            for args in commands:
                with self.subTest(command=args[0]):
                    result = subprocess.run([sys.executable, str(ROOT/args[0]), *args[1:]],
                                            cwd=ROOT, capture_output=True, text=True)
                    self.assertEqual(result.returncode, 0, result.stdout+result.stderr)
            self.assertTrue(list(work.rglob('*.xlsx')))
            self.assertTrue(list(work.rglob('*.html')))


if __name__ == '__main__':
    unittest.main()
