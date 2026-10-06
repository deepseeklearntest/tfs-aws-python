#!/usr/bin/env python3
"""
AWS Proactive Excel Report Generator

Reads the latest findings CSVs from the findings directory and produces
a manager-ready Excel workbook with RAG-coloured sheets per wave.

Run: python generate_report.py
     python generate_report.py --findings-dir findings --output-dir reports
"""
from __future__ import annotations

import argparse
import csv
import logging
import re
import sys
from datetime import datetime
from pathlib import Path

from common.findings_io import WAVE_MAP, severity_of, discover_latest_csvs

from common.spreadsheet import append_literal_row, set_literal_cell

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

WAVE_LABELS: dict[str, str] = {
    'wave1_cost':         'Wave 1 – Cost',
    'wave2_lifecycle':    'Wave 2 – Lifecycle',
    'wave3_reliability':  'Wave 3 – Reliability',
    'wave4_security':     'Wave 4 – Security',
    'wave5_performance':  'Wave 5 – Performance',
    'continuous':         'Continuous',
}

TAB_COLORS: dict[str, str] = {
    'Executive Summary':            '1F4E79',
    'Top Savings':                  '375623',
    'By Account':                   '595959',
    'Wave 1 – Cost':          '008080',
    'Wave 2 – Lifecycle':     '7030A0',
    'Wave 3 – Reliability':   'FF6600',
    'Wave 4 – Security':      'C00000',
    'Wave 5 – Performance':   '4472C4',
    'Continuous':                   '2E75B6',
    'All Findings':                 '595959',
}

# Standard columns written to every data sheet (wave sheets, all-findings)
STANDARD_COLS = [
    'account_id', 'region', 'resource_id', 'resource_type',
    'finding', 'detail',
    'estimated_monthly_savings_usd', 'estimated_annual_savings_usd',
    'source_script', 'tags', 'owner',
]

# Columns for Top Savings sheet (no resource_id per spec)
TOP_SAVINGS_COLS = [
    'account_id', 'region', 'resource_type', 'finding', 'detail',
    'estimated_monthly_savings_usd', 'estimated_annual_savings_usd', 'source_script',
]

# Columns for By Account sheet
BY_ACCOUNT_COLS = [
    'account_id', 'total_findings',
    'est_monthly_savings_usd', 'est_annual_savings_usd',
    'wave1_cost_findings', 'wave2_lifecycle_findings',
    'wave3_reliability_findings', 'wave4_security_findings',
    'wave5_performance_findings',
]

# ---------------------------------------------------------------------------
# RAG palette
# ---------------------------------------------------------------------------
_RED_FILL   = PatternFill('solid', fgColor='FFC7CE')
_RED_FONT   = Font(color='9C0006')
_AMBER_FILL = PatternFill('solid', fgColor='FFEB9C')
_AMBER_FONT = Font(color='9C6500')
_GREEN_FILL = PatternFill('solid', fgColor='C6EFCE')
_GREEN_FONT = Font(color='276221')

_HEADER_FILL = PatternFill('solid', fgColor='1F4E79')
_HEADER_FONT = Font(bold=True, color='FFFFFF')

# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------

def discover_csvs(findings_dir: str) -> list[Path]:
    return discover_latest_csvs(findings_dir)


def wave_of(script_name: str) -> str | None:
    """Map script name (e.g. '06_detect_nat_gateway_hotpaths') to a wave key."""
    m = re.match(r'^(\d+)', script_name)
    if not m:
        return None
    prefix = m.group(1).zfill(2)
    return WAVE_MAP.get(prefix)


def _safe_float(value) -> float:
    """Convert any CSV cell to float; return 0.0 on blank or non-numeric."""
    try:
        return float(value) if value else 0.0
    except (ValueError, TypeError):
        return 0.0


def _normalize_account_id(value: str) -> str:
    """Convert scientific-notation account IDs (e.g. '8.60118E+11') to decimal strings.

    This happens when a CSV is opened and re-saved by Excel, which converts
    12-digit account IDs to scientific notation, losing the last ~6 digits.
    """
    if not value or ('E' not in value.upper()):
        return value
    try:
        return str(int(float(value)))
    except (ValueError, TypeError):
        return value


def load_findings(csv_paths: list[Path]) -> list[dict]:
    """Read all CSVs into a list of dicts; add source_script and annual savings columns."""
    findings: list[dict] = []
    for csv_path in csv_paths:
        m = re.match(r'^(.+?)_(\d{8})_(\d{6})$', csv_path.stem)
        script_name = m.group(1) if m else csv_path.stem
        try:
            with open(csv_path, newline='', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                row_count = 0
                for row in reader:
                    try:
                        row['account_id'] = _normalize_account_id(row.get('account_id', ''))
                        monthly = _safe_float(row.get('estimated_monthly_savings_usd'))
                        row['estimated_monthly_savings_usd'] = monthly
                        row['estimated_annual_savings_usd'] = round(monthly * 12, 2)
                        row['source_script'] = script_name
                        findings.append(row)
                        row_count += 1
                    except Exception as exc:
                        log.warning("Skipping malformed row in %s: %s", csv_path.name, exc)
            print(f"  Loaded: {script_name:<55} ({row_count} rows)")
        except Exception as exc:
            log.warning("Could not read %s: %s", csv_path, exc)
    return findings


def load_from_inventory(findings_dir: str) -> list[dict]:
    """Load findings from the latest inventory_*.csv written by 54_inventory_refresh.py."""
    root = Path(findings_dir)
    candidates = [
        p for p in root.glob('inventory_*.csv')
        if re.match(r'^inventory_\d{8}_\d{6}\.csv$', p.name)
    ]
    if not candidates:
        return []
    inventory_path = max(candidates, key=lambda p: p.name)
    findings: list[dict] = []
    try:
        with open(inventory_path, newline='', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                try:
                    row['account_id'] = _normalize_account_id(row.get('account_id', ''))
                    row['estimated_monthly_savings_usd'] = _safe_float(row.get('estimated_monthly_savings_usd'))
                    row['estimated_annual_savings_usd'] = _safe_float(row.get('estimated_annual_savings_usd'))
                    findings.append(row)
                except Exception as exc:
                    log.warning('Skipping malformed row in %s: %s', inventory_path.name, exc)
        print(f'  Loaded inventory: {inventory_path.name} ({len(findings)} rows)')
    except Exception as exc:
        log.warning('Could not read %s: %s', inventory_path, exc)
    return findings


# ---------------------------------------------------------------------------
# RAG helpers — uniform signature: (finding, monthly_usd) -> (fill, font)
# ---------------------------------------------------------------------------

RagResult = tuple[PatternFill, Font]


def _rag(wave: str, finding: str, monthly_usd: float) -> RagResult:
    return {
        'High': (_RED_FILL, _RED_FONT),
        'Medium': (_AMBER_FILL, _AMBER_FONT),
        'Low': (_GREEN_FILL, _GREEN_FONT),
    }[severity_of(wave, finding, monthly_usd)]


def rag_cost(finding: str, monthly_usd: float) -> RagResult:
    return _rag('wave1_cost', finding, monthly_usd)


def rag_lifecycle(finding: str, monthly_usd: float) -> RagResult:
    return _rag('wave2_lifecycle', finding, monthly_usd)


def rag_reliability(finding: str, monthly_usd: float) -> RagResult:
    return _rag('wave3_reliability', finding, monthly_usd)


def rag_security(finding: str, monthly_usd: float) -> RagResult:
    return _rag('wave4_security', finding, monthly_usd)


def rag_continuous(finding: str, monthly_usd: float) -> RagResult:
    return _rag('continuous', finding, monthly_usd)


# ---------------------------------------------------------------------------
# Sheet formatting
# ---------------------------------------------------------------------------

def style_sheet(ws) -> None:
    """Dark-blue header row, freeze pane at A2, auto-fit columns (min 12, max 60)."""
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=False)

    ws.freeze_panes = 'A2'

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(
            (len(str(cell.value)) for cell in col if cell.value is not None),
            default=0,
        )
        ws.column_dimensions[col_letter].width = max(12, min(60, max_len + 2))
        if col[0].value == 'account_id':
            for cell in col[1:]:
                cell.number_format = '@'


# ---------------------------------------------------------------------------
# Sheet writers
# ---------------------------------------------------------------------------

def write_top_savings_sheet(wb: Workbook, findings: list[dict]) -> None:
    """All findings with monthly savings > $0, sorted descending, RAG by savings tier."""
    savings_rows = [
        r for r in findings
        if _safe_float(r.get('estimated_monthly_savings_usd')) > 0
    ]
    if not savings_rows:
        return

    savings_rows.sort(
        key=lambda r: _safe_float(r.get('estimated_monthly_savings_usd')),
        reverse=True,
    )

    ws = wb.create_sheet('Top Savings')
    ws.append(TOP_SAVINGS_COLS)

    for row in savings_rows:
        monthly = _safe_float(row.get('estimated_monthly_savings_usd'))
        fill, font = rag_cost('', monthly)
        append_literal_row(ws, [row.get(c, '') for c in TOP_SAVINGS_COLS])
        for cell in ws[ws.max_row]:
            cell.fill = fill
            cell.font = font

    m_idx = TOP_SAVINGS_COLS.index('estimated_monthly_savings_usd') + 1
    a_idx = TOP_SAVINGS_COLS.index('estimated_annual_savings_usd') + 1
    for data_row in ws.iter_rows(min_row=2):
        data_row[m_idx - 1].number_format = '$#,##0.00'
        data_row[a_idx - 1].number_format = '$#,##0.00'

    style_sheet(ws)


def write_by_account_sheet(wb: Workbook, findings: list[dict]) -> None:
    """One summary row per account_id with per-wave finding counts and total savings."""
    accounts: dict[str, dict] = {}

    for row in findings:
        acct = row.get('account_id', 'unknown')
        if acct not in accounts:
            accounts[acct] = {
                'account_id': acct,
                'total_findings': 0,
                'est_monthly_savings_usd': 0.0,
                'est_annual_savings_usd': 0.0,
                'wave1_cost_findings': 0,
                'wave2_lifecycle_findings': 0,
                'wave3_reliability_findings': 0,
                'wave4_security_findings': 0,
                'wave5_performance_findings': 0,
            }
        a = accounts[acct]
        a['total_findings'] += 1
        monthly = _safe_float(row.get('estimated_monthly_savings_usd'))
        a['est_monthly_savings_usd'] += monthly
        a['est_annual_savings_usd'] += monthly * 12

        wave = wave_of(row.get('source_script', ''))
        wave_col = f'{wave}_findings' if wave else None
        if wave_col and wave_col in a:
            a[wave_col] += 1

    if not accounts:
        return

    ws = wb.create_sheet('By Account')
    ws.append(BY_ACCOUNT_COLS)

    for a in sorted(accounts.values(), key=lambda x: x['account_id']):
        a['est_monthly_savings_usd'] = round(a['est_monthly_savings_usd'], 2)
        a['est_annual_savings_usd'] = round(a['est_annual_savings_usd'], 2)
        append_literal_row(ws, [a.get(c, 0) for c in BY_ACCOUNT_COLS])

    m_idx = BY_ACCOUNT_COLS.index('est_monthly_savings_usd') + 1
    a_idx = BY_ACCOUNT_COLS.index('est_annual_savings_usd') + 1
    for data_row in ws.iter_rows(min_row=2):
        data_row[m_idx - 1].number_format = '$#,##0.00'
        data_row[a_idx - 1].number_format = '$#,##0.00'

    style_sheet(ws)


def write_wave_sheet(
    wb: Workbook,
    findings: list[dict],
    wave: str,
    label: str,
    rag_fn,
) -> None:
    """Generic wave sheet. Filters by wave, adds extra columns found in data, applies rag_fn."""
    wave_rows = [r for r in findings if wave_of(r.get('source_script', '')) == wave]
    if not wave_rows:
        return

    # Include script-specific extra columns that appear in this wave's data
    extra_cols: list[str] = []
    for row in wave_rows:
        for key in row:
            if key not in STANDARD_COLS and key not in extra_cols:
                extra_cols.append(key)

    cols = STANDARD_COLS + extra_cols

    ws = wb.create_sheet(label)
    ws.append(cols)

    for row in wave_rows:
        monthly = _safe_float(row.get('estimated_monthly_savings_usd'))
        finding = row.get('finding', '')
        fill, font = rag_fn(finding, monthly)
        append_literal_row(ws, [row.get(c, '') for c in cols])
        for cell in ws[ws.max_row]:
            cell.fill = fill
            cell.font = font

    m_idx = cols.index('estimated_monthly_savings_usd') + 1
    a_idx = cols.index('estimated_annual_savings_usd') + 1
    for data_row in ws.iter_rows(min_row=2):
        data_row[m_idx - 1].number_format = '$#,##0.00'
        data_row[a_idx - 1].number_format = '$#,##0.00'

    style_sheet(ws)


def write_summary_sheet(wb: Workbook, findings: list[dict]) -> None:
    """Executive Summary: title, key metrics box, wave breakdown, top 10 action items."""
    ws = wb.create_sheet('Executive Summary')

    # --- Title ---
    ws['A1'] = 'AWS Proactive Findings Report'
    ws['A1'].font = Font(bold=True, size=16, color='1F4E79')

    accounts = sorted({f.get('account_id', '') for f in findings if f.get('account_id')})
    scripts  = sorted({f.get('source_script', '') for f in findings if f.get('source_script')})
    now_str  = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    ws['A2'] = (
        f"Generated: {now_str}  |  Accounts scanned: {len(accounts)}"
        f"  |  Scripts run: {len(scripts)}"
    )

    # --- Key Metrics ---
    ws['A4'] = 'KEY METRICS'
    ws['A4'].font = Font(bold=True)

    for col, header in enumerate(['Total Findings', 'Monthly Saves', 'Annual Saves'], 1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    total_monthly = sum(_safe_float(f.get('estimated_monthly_savings_usd')) for f in findings)
    ws.cell(row=6, column=1, value=len(findings))
    m_cell = ws.cell(row=6, column=2, value=round(total_monthly, 2))
    a_cell = ws.cell(row=6, column=3, value=round(total_monthly * 12, 2))
    m_cell.number_format = '$#,##0.00'
    a_cell.number_format = '$#,##0.00'

    # --- Wave Breakdown ---
    ws['A8'] = 'BREAKDOWN BY WAVE'
    ws['A8'].font = Font(bold=True)

    for col, header in enumerate(['Wave', 'Findings', 'Monthly $', 'Annual $'], 1):
        cell = ws.cell(row=9, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    ZERO_SAVINGS_WAVES = {'wave3_reliability', 'wave4_security'}
    wave_order = [
        'wave1_cost', 'wave2_lifecycle', 'wave3_reliability',
        'wave4_security', 'wave5_performance', 'continuous',
    ]
    data_row = 10
    for wave_key in wave_order:
        wave_rows = [f for f in findings if wave_of(f.get('source_script', '')) == wave_key]
        if not wave_rows:
            continue
        monthly = (
            0.0 if wave_key in ZERO_SAVINGS_WAVES
            else sum(_safe_float(r.get('estimated_monthly_savings_usd')) for r in wave_rows)
        )
        ws.cell(row=data_row, column=1, value=WAVE_LABELS[wave_key])
        ws.cell(row=data_row, column=2, value=len(wave_rows))
        m = ws.cell(row=data_row, column=3, value=round(monthly, 2))
        a = ws.cell(row=data_row, column=4, value=round(monthly * 12, 2))
        m.number_format = '$#,##0.00'
        a.number_format = '$#,##0.00'
        data_row += 1

    data_row += 1
    note = ws.cell(
        row=data_row, column=1,
        value=(
            'Note: Security and Reliability findings carry $0 estimated savings. '
            'They are risk items requiring remediation, not direct cost savings.'
        ),
    )
    note.font = Font(italic=True, color='595959')

    # --- Top 10 Action Items ---
    data_row += 2
    ws.cell(row=data_row, column=1, value='TOP 10 ACTION ITEMS  (by monthly savings)').font = Font(bold=True)
    data_row += 1

    for col, header in enumerate(['Rank', 'Account', 'Resource Type', 'Finding', 'Monthly $', 'Annual $'], 1):
        cell = ws.cell(row=data_row, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    data_row += 1

    top10 = sorted(
        findings,
        key=lambda r: _safe_float(r.get('estimated_monthly_savings_usd')),
        reverse=True,
    )[:10]

    for rank, row in enumerate(top10, 1):
        monthly = _safe_float(row.get('estimated_monthly_savings_usd'))
        ws.cell(row=data_row, column=1, value=rank)
        acct_cell = set_literal_cell(ws, data_row, 2, row.get('account_id', ''))
        acct_cell.number_format = '@'
        set_literal_cell(ws, data_row, 3, row.get('resource_type', ''))
        set_literal_cell(ws, data_row, 4, row.get('finding', ''))
        m = ws.cell(row=data_row, column=5, value=round(monthly, 2))
        a = ws.cell(row=data_row, column=6, value=round(monthly * 12, 2))
        m.number_format = '$#,##0.00'
        a.number_format = '$#,##0.00'
        data_row += 1

    # Fixed column widths -- non-standard layout so style_sheet is not used
    ws.column_dimensions['A'].width = 55
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 50
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15


def write_all_findings_sheet(wb: Workbook, findings: list[dict]) -> None:
    """All findings, standard columns only, sorted by monthly savings descending. No RAG."""
    sorted_rows = sorted(
        findings,
        key=lambda r: _safe_float(r.get('estimated_monthly_savings_usd')),
        reverse=True,
    )

    ws = wb.create_sheet('All Findings')
    ws.append(STANDARD_COLS)

    for row in sorted_rows:
        append_literal_row(ws, [row.get(c, '') for c in STANDARD_COLS])

    m_idx = STANDARD_COLS.index('estimated_monthly_savings_usd') + 1
    a_idx = STANDARD_COLS.index('estimated_annual_savings_usd') + 1
    for data_row in ws.iter_rows(min_row=2):
        data_row[m_idx - 1].number_format = '$#,##0.00'
        data_row[a_idx - 1].number_format = '$#,##0.00'

    style_sheet(ws)


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

WAVE_CONFIG = [
    ('wave1_cost',        WAVE_LABELS['wave1_cost'],        rag_cost),
    ('wave2_lifecycle',   WAVE_LABELS['wave2_lifecycle'],   rag_lifecycle),
    ('wave3_reliability', WAVE_LABELS['wave3_reliability'], rag_reliability),
    ('wave4_security',    WAVE_LABELS['wave4_security'],    rag_security),
    ('wave5_performance', WAVE_LABELS['wave5_performance'], rag_cost),
    ('continuous',        WAVE_LABELS['continuous'],        rag_continuous),
]


def main() -> None:
    logging.basicConfig(level=logging.WARNING, format='%(levelname)s %(message)s')

    parser = argparse.ArgumentParser(
        description='Generate Excel report from AWS Proactive findings CSVs.'
    )
    parser.add_argument(
        '--findings-dir', default='findings',
        help='Root directory to scan for CSV files (default: findings)',
    )
    parser.add_argument(
        '--output-dir', default='reports',
        help='Directory to write the Excel file (default: reports)',
    )
    parser.add_argument(
        '--from-inventory', action='store_true',
        help='Read from the latest inventory_*.csv instead of re-scanning individual CSVs',
    )
    args = parser.parse_args()

    if args.from_inventory:
        findings = load_from_inventory(args.findings_dir)
        if not findings:
            print(
                f"No inventory_*.csv found in {args.findings_dir}/ — "
                "run continuous/54_inventory_refresh.py first or omit --from-inventory",
                file=sys.stderr,
            )
            sys.exit(1)
    else:
        print(f"Scanning {args.findings_dir}/ for CSVs...")
        try:
            csv_paths = discover_csvs(args.findings_dir)
        except FileNotFoundError as exc:
            print(f"Error: {exc}", file=sys.stderr)
            sys.exit(1)
        if not csv_paths:
            print(
                f"No CSVs found in {args.findings_dir}/ "
                "(searched recursively, excluding logs/ and inventory_*)",
                file=sys.stderr,
            )
            sys.exit(1)
        findings = load_findings(csv_paths)

    accounts = {f.get('account_id', '') for f in findings}
    print(f"\nTotal: {len(findings)} findings across {len(accounts)} accounts")

    print("\nWriting report...")
    wb = Workbook()
    wb.remove(wb.active)

    write_summary_sheet(wb, findings)
    print("  + Executive Summary")

    write_top_savings_sheet(wb, findings)
    if 'Top Savings' in {ws.title for ws in wb.worksheets}:
        savings_count = sum(1 for f in findings if _safe_float(f.get('estimated_monthly_savings_usd')) > 0)
        print(f"  + Top Savings ({savings_count} rows with savings > $0)")

    write_by_account_sheet(wb, findings)
    if 'By Account' in {ws.title for ws in wb.worksheets}:
        print(f"  + By Account ({len(accounts)} accounts)")

    for wave_key, label, rag_fn in WAVE_CONFIG:
        before = {ws.title for ws in wb.worksheets}
        write_wave_sheet(wb, findings, wave_key, label, rag_fn)
        if label not in before and label in {ws.title for ws in wb.worksheets}:
            wave_count = sum(1 for f in findings if wave_of(f.get('source_script', '')) == wave_key)
            print(f"  + {label} ({wave_count} rows)")

    write_all_findings_sheet(wb, findings)
    print(f"  + All Findings ({len(findings)} rows)")

    for ws in wb.worksheets:
        color = TAB_COLORS.get(ws.title)
        if color:
            ws.sheet_properties.tabColor = color

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_path = output_dir / f"aws_proactive_report_{timestamp}.xlsx"
    wb.save(str(output_path))
    print(f"\nReport saved: {output_path}")


if __name__ == '__main__':
    main()
